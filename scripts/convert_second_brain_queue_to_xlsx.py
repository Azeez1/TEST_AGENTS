import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


MD_PATH = Path(
    r"C:\Users\sabaa\OneDrive\Desktop\MEMORY\VAULT\wiki\queries\Second Brain Review Queue 2026-06-30.md"
)
OUT_PATH = Path(
    r"C:\Users\sabaa\OneDrive\Desktop\MEMORY\VAULT\wiki\queries\Second Brain Review Queue 2026-06-30.xlsx"
)


def section(lines, name):
    start = None
    for i, line in enumerate(lines):
        if line.strip() == f"## {name}":
            start = i
            break
    if start is None:
        return []
    end = len(lines)
    for j in range(start + 1, len(lines)):
        if lines[j].startswith("## "):
            end = j
            break
    return lines[start:end]


def split_md_row(line):
    row = line.strip()
    if not row.startswith("|"):
        return []
    return [c.strip().replace("`", "") for c in row.strip("|").split("|")]


def parse_table(sec_lines):
    header = None
    rows = []
    for line in sec_lines:
        if not line.startswith("|"):
            continue
        cells = split_md_row(line)
        if not cells:
            continue
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            continue
        if header is None:
            header = cells
        else:
            rows.append(cells)
    return header or [], rows


def normalize(header, rows):
    clean = []
    for row in rows:
        clean.append(row[: len(header)] + [""] * max(0, len(header) - len(row)))
    return clean


def add_table(ws, sheet_name):
    if ws.max_row < 2 or ws.max_column < 1:
        return
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    display = re.sub(r"[^A-Za-z0-9_]", "", sheet_name)[:20] + "Table"
    tab = Table(displayName=display, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(tab)


def set_widths(ws, header):
    for col in range(1, ws.max_column + 1):
        name = str(ws.cell(1, col).value or "").lower()
        if name in {"review", "suggested", "priority", "score", "msgs", "size", "date"}:
            width = 14
        elif name in {"url", "path"}:
            width = 42
        elif name in {
            "why it might matter",
            "suggested next action",
            "guardrail",
            "existing summary / last seen",
            "first prompt",
            "note",
        }:
            width = 48
        elif name in {"source", "file", "session id"}:
            width = 34
        else:
            width = 24
        ws.column_dimensions[get_column_letter(col)].width = width


def main():
    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    section_map = {
        "High-Value Sessions": "High-Value Claude Session Candidates",
        "Maybe Sessions": "Maybe Claude Session Candidates",
        "Skip Sessions": "Skip / Parking Lot Claude Sessions",
        "Vault Intake": "Vault Intake Sources",
        "New Knowledge": "Research-Backed New Knowledge Candidates",
        "Manual Intake": "New Knowledge Intake Lane",
    }

    parsed = {}
    for sheet, sec in section_map.items():
        header, rows = parse_table(section(lines, sec))
        parsed[sheet] = (header, normalize(header, rows))

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    summary_rows = [
        ["Review Queue Workbook", "Second Brain Review Queue 2026-06-30"],
        ["Source Markdown", str(MD_PATH)],
        ["Canonical Ingest", "None. This workbook is review-only."],
        [],
        ["Tab", "Rows", "What to do"],
    ]
    for sheet, (_, rows) in parsed.items():
        if sheet == "Manual Intake":
            note = "Add anything new you personally want reviewed."
        elif sheet == "New Knowledge":
            note = "Approve sources worth scraping, reading, or adding to raw notes."
        elif "Sessions" in sheet:
            note = "Approve sessions worth distilling into clean raw notes."
        else:
            note = "Approve files worth duplicate-checking and promoting."
        summary_rows.append([sheet, len(rows), note])
    summary_rows.extend(
        [
            [],
            ["Decision Guide", ""],
            ["APPROVE", "Worth processing into raw notes or ingestion candidates."],
            ["MAYBE", "Needs quick skim before deciding."],
            ["SKIP", "Do not process unless you change your mind."],
        ]
    )
    for row in summary_rows:
        summary.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(bottom=thin)
    linkish = re.compile(r"^(https?://|[A-Za-z]:\\)")

    for cell in summary[1]:
        cell.font = Font(bold=True, size=14)
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 90
    summary.column_dimensions["C"].width = 70
    summary.freeze_panes = "A6"

    for sheet, (header, rows) in parsed.items():
        ws = wb.create_sheet(sheet[:31])
        if not header:
            ws.append(["No data found"])
            continue
        ws.append(header)
        for row in rows:
            ws.append(row)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and linkish.match(cell.value):
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"

        set_widths(ws, header)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        add_table(ws, sheet)

        if "Review" in header:
            col = header.index("Review") + 1
            for r in range(2, ws.max_row + 1):
                ws.cell(r, col).value = ""
            dv = DataValidation(type="list", formula1='"APPROVE,MAYBE,SKIP"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{ws.max_row}")

        if "Decision" in header:
            col = header.index("Decision") + 1
            dv = DataValidation(type="list", formula1='"APPROVE,MAYBE,SKIP"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{ws.max_row}")

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = border

    wb.save(OUT_PATH)

    check = load_workbook(OUT_PATH, read_only=True)
    print(OUT_PATH)
    for name in check.sheetnames:
        ws = check[name]
        rows = ws.max_row if name == "Summary" else max(0, ws.max_row - 1)
        print(f"{name}: {rows} rows")


if __name__ == "__main__":
    main()
